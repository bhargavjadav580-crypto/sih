"""Recreate six editable-content slides on the user's exact PDF template geometry.
Template decoration is preserved as raster artwork. All new content and diagrams
are editable PowerPoint objects. No official branding approval is implied.
"""
from pathlib import Path
from PIL import Image, ImageOps
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
import json

ROOT=Path('/app/presentation'); OUT=Path('/app/frontend/public/downloads'); OUT.mkdir(parents=True,exist_ok=True)
W,H=792,612
BLUE='234E81'; FOOTER='0875B9'; INK='263634'; GREEN='0B6E4F'; GRAY='66756E'; WHITE='FFFFFF'; LIGHT='EDF3F7'
prs=Presentation(); prs.slide_width=Pt(W); prs.slide_height=Pt(H)
prs.core_properties.title='LABOURLINK — Cooperative Gig Services Platform'
prs.core_properties.subject='SIH PS 26089 · six-slide synthetic-prototype presentation'
prs.core_properties.author='LABOURLINK team · AI-assisted preparation'
prs.core_properties.comments='PDF template recreated at 792 × 612 pt. Team registration and official branding approval require confirmation.'
manifest=[]

def shape(slide,x,y,w,h,fill=WHITE,line=None,kind=MSO_SHAPE.RECTANGLE):
 s=slide.shapes.add_shape(kind,Pt(x),Pt(y),Pt(w),Pt(h));s.fill.solid();s.fill.fore_color.rgb=RGBColor.from_string(fill)
 style=s._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
 if style is not None:s._element.remove(style)
 s._element.spPr.append(OxmlElement('a:effectLst'))
 if line:s.line.color.rgb=RGBColor.from_string(line);s.line.width=Pt(.7)
 else:s.line.fill.background()
 return s

def text(slide,x,y,w,h,content,size=12,color=INK,bold=False,align=None,font='Arial',url=None):
 box=slide.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(h));tf=box.text_frame;tf.clear();tf.word_wrap=True
 tf.margin_left=tf.margin_right=Pt(0);tf.margin_top=tf.margin_bottom=Pt(0)
 for i,line in enumerate(content.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color);p.space_after=Pt(2);p.space_before=Pt(0);p.line_spacing=1.08
  if align:p.alignment=align
  if url:
   for run in p.runs:run.hyperlink.address=url;run.font.underline=True
 manifest.append({'slide':len(prs.slides),'x':x,'y':y,'w':w,'h':h,'text':content,'fontSize':size,'url':url})
 return box

def bullet(slide,x,y,w,content,size=12):
 text(slide,x,y,9,20,'•',size,BLUE,bold=True);return text(slide,x+14,y,w-14,33,content,size)

def header(slide,x,y,w,content,size=13):return text(slide,x,y,w,35,content,size,BLUE,True)

def connector(slide,x1,y1,x2,y2):
 c=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Pt(x1),Pt(y1),Pt(x2),Pt(y2));c.line.color.rgb=RGBColor.from_string('527D97');c.line.width=Pt(1.25)
 if y2>y1:
  end=OxmlElement('a:tailEnd');end.set('type','triangle');end.set('w','sm');end.set('len','sm');c.line._get_or_add_ln().append(end)
 return c

def node(slide,x,y,w,h,label,fill=LIGHT,size=10):
 shape(slide,x,y,w,h,fill,'C9D8E1',MSO_SHAPE.ROUNDED_RECTANGLE)
 return text(slide,x+5,y+6,w-10,h-7,label,size,BLUE,True,PP_ALIGN.CENTER)

def contain(slide,path,x,y,w,h):
 im=Image.open(path);iw,ih=im.size;scale=min(w/iw,h/ih);nw,nh=iw*scale,ih*scale
 return slide.shapes.add_picture(str(path),Pt(x+(w-nw)/2),Pt(y+(h-nh)/2),width=Pt(nw),height=Pt(nh))

def prepare_logo():
 src=Image.open('/app/source-assets/sih-2026-user-supplied.png').convert('RGB')
 # User-provided bitmap is already labeled 2026. Crop and rearrange without changing its year.
 sx,sy=src.width/622,src.height/802
 icon=src.crop(tuple(int(v) for v in (120*sx,82*sy,432*sx,447*sy)))
 words=Image.new('RGB',(500,230),'white')
 for bounds,y,target in [((86,497,480,553),0,(420,65)),((86,574,544,631),77,(485,66)),((193,652,373,715),154,(192,68))]:
  crop=src.crop(tuple(int(v*(sx if i%2==0 else sy)) for i,v in enumerate(bounds)))
  crop=ImageOps.contain(crop,target);words.paste(crop,((500-crop.width)//2,y))
 canvas=Image.new('RGB',(600,290),'white')
 icon=ImageOps.contain(icon,(223,263));words=ImageOps.contain(words,(360,200))
 canvas.paste(icon,(0,(290-icon.height)//2));canvas.paste(words,(231,(290-words.height)//2))
 path=ROOT/'sih-2026-lockup.png';canvas.save(path);return path

logo=prepare_logo()
def base(index):
 slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.shapes.add_picture(str(ROOT/f'template-{index}.png'),0,0,width=Pt(W),height=Pt(H))
 # Exact original top-right logo area, old 2025 lockup removed.
 shape(slide,641,87,150,75);contain(slide,logo,643,90,141,66)
 if index>1:
  # Keep oval size and position; remove the placeholder by recreating the same outline.
  shape(slide,17,95,90,61)
  shape(slide,21,100,82,52,WHITE,'8060A8',MSO_SHAPE.OVAL)
  text(slide,28,114,68,28,'LABOURLINK',9.5,INK,True,PP_ALIGN.CENTER)
  # Content area only. Header, footer, page number and their positions are preserved.
  shape(slide,0,200,792,296)
 else:
  shape(slide,85,96,548,43)
  text(slide,96,99,530,39,'SMART INDIA HACKATHON 2026',29,BLUE,True,font='Times New Roman')
  shape(slide,15,248,347,303)
  shape(slide,350,489,29,47)
 return slide

# 1 — fixed title fields; original right-hand graphic remains untouched.
s=base(1)
fields=[('Problem Statement ID','26089'),('Problem Statement Title','Cooperative Gig Services Platform for\nHousehold & Community Services'),('Theme','Agriculture, FoodTech & Rural Development'),('PS Category','Software'),('Team ID','[REGISTERED TEAM ID]'),('Team Name (Registered on portal)','LABOURLINK [confirm portal spelling]')]
for (label,value),y in zip(fields,[254,301,362,410,455,502]):
 text(s,29,y,326,13,label.upper(),9.5,BLUE,True)
 text(s,29,y+15,330,33,value,13.4 if len(value)>40 else 15,INK,True)
text(s,30,556,720,17,'Connecting Communities with Verified Cooperative Workers.',12,GREEN,True)
text(s,30,579,725,14,'User-supplied 2026 branding · official approval and team details require confirmation.',9.2,GRAY)
s.notes_slide.notes_text_frame.text='Problem statement title and theme are taken from the user-provided organizer statement. Do not substitute unofficial themes. Team ID and exact portal spelling remain unconfirmed. The 2026 image is user-supplied, not independently verified as an approved event asset. AI-assisted preparation does not meet any wholly human-created requirement.'

# 2 — prescribed pointers and one editable coordination diagram.
s=base(2)
text(s,286,169,220,25,'LABOURLINK',18,GREEN,True,PP_ALIGN.CENTER)
header(s,29,223,745,'Proposed Solution (Describe your Idea/Solution/Prototype)',18)
header(s,29,263,440,'Detailed explanation of the proposed solution',12)
bullet(s,31,284,421,'Societies import, verify and manage their worker rosters.',11.5)
bullet(s,31,303,421,'Households book services; institutions request multiple workers.',11.5)
bullet(s,31,322,421,'Workers review terms, accept work and track earnings.',11.5)
header(s,29,350,430,'How it addresses the problem',12)
text(s,43,370,420,39,'Society records → verified profiles → suitable matches\n→ completed services → transparent settlement',11.5)
header(s,29,409,430,'Innovation and uniqueness of the solution',12)
bullet(s,31,429,421,'Society-controlled verification and approved policies.',11)
bullet(s,31,446,421,'Separate payment, earning and settlement records.',11)
bullet(s,31,463,421,'Capacity-aware coordination across participating societies.',11)
node(s,527,265,184,31,'Customer requirement',size=11)
connector(s,619,296,619,310);text(s,614,296,16,16,'↓',11,BLUE)
node(s,527,312,184,31,'Chosen / suitable society',size=11)
connector(s,619,343,619,355);connector(s,518,355,718,355);connector(s,518,355,518,368);connector(s,718,355,718,368)
node(s,466,369,125,32,'Capacity available',size=10.6)
node(s,630,369,133,32,'Shortfall',size=10.6)
connector(s,528,401,528,415);connector(s,696,401,696,415)
node(s,466,417,125,46,'Worker offers\n→ acceptance',fill='EEF4EA',size=10.8)
node(s,630,417,133,46,'Authorized federation\n→ other societies',size=10)
text(s,466,471,300,21,'Affiliation + authority; no mandatory middleman.',9.2,GRAY)
s.notes_slide.notes_text_frame.text='Federation involvement depends on affiliation, authorized scope and customer consent, not a fixed headcount threshold. Capacity is filtered by eligible skills and available time before ranking. A federation can contract or execute work where its own rules authorize it; this diagram is not a universal legal hierarchy or payment chain.'

# 3 — technical boundaries and real prototype screenshots, if captured.
s=base(3)
text(s,31,181,730,18,'Six experiences: Household · Institution · Worker · Society · District · National',11.5,BLUE,True)
header(s,31,212,354,'Technologies to be used',14)
text(s,31,232,350,20,'(programming languages, frameworks, hardware)',9.4,GRAY)
bullet(s,31,254,350,'Frontend: React + TypeScript; EN / HI / GU.',11.5)
bullet(s,31,276,350,'Planned backend: FastAPI.',11.5)
bullet(s,31,298,350,'Planned database: MongoDB + geospatial indexing.',11.5)
bullet(s,31,320,350,'Future authorized providers: payments, OTP, maps, voice, welfare.',11.5)
header(s,419,212,344,'Methodology and process for implementation',13)
text(s,419,233,344,17,'(Flow Charts/Images/ working prototype)',9.4,GRAY)
text(s,419,254,344,43,'Import → verify → match → book → accept\n→ complete → record payment → reconcile',12,INK,True)
text(s,419,301,344,32,'Matching: skill + availability first; then distance, experience, rating and workload.',10.5)
text(s,419,335,344,32,'Forecast: history → demand baseline → capacity gap → admin review.',10.5)
for title,path,x in [('Worker verification',ROOT/'verification-crop.jpg',31),('Separate settlement ledger',ROOT/'settlement-crop.jpg',410)]:
 text(s,x,372,350,15,title+' · synthetic prototype',10.5,BLUE,True)
 if path.exists():contain(s,path,x,391,350,63)
 else:
  shape(s,x,391,350,63,'F2F5F7','D3DEE5');text(s,x+12,412,326,30,'Proposed interface — screenshot pending',11,GRAY,False,PP_ALIGN.CENTER)
shape(s,26,463,740,26,'EAF0F5')
text(s,35,468,722,22,'Prototype scope: synthetic data and simulated handoffs. No live payment, government-data connection, backend or production AI claim.',9.8,BLUE,True)
s.notes_slide.notes_text_frame.text='Six frontend experiences: Household, Institution, Worker, Society, District and National. The current preview is one codebase with distinct routes and shared synthetic state; it is not six independently verified deployments. Separate origins need explicit schema-validated export/import. Browser role views are not secure authorization. Production integrations remain future work. Screenshots are actual local synthetic prototype screens, when present.'

# 4 — feasibility, four risks, pilot and production gate.
s=base(4)
header(s,31,211,730,'Analysis of the feasibility of the idea',14)
bullet(s,31,236,730,'Start with participating societies’ own records—not assumed government worker-data access.',12)
bullet(s,31,258,730,'Begin with one society and selected household/community services.',12)
bullet(s,31,280,730,'Agree cooperative operating charges; do not assume a universal commission.',12)
header(s,31,311,335,'Potential challenges and risks',12.5)
header(s,389,311,372,'Strategies for overcoming these challenges',12.5)
rows=[('Incomplete worker records','Validated imports, review queue, verification history'),('Language / digital-literacy barriers','Simple multilingual screens and assisted onboarding'),('Overbooking / bulk shortfalls','Time checks, reservations, partial-fulfilment review'),('Cash, no-shows, payment disputes','Separate ledgers, evidence trail, society review')]
for i,(risk,response) in enumerate(rows):
 y=337+i*26;shape(s,30,y,731,25,'F0F4F8' if i%2==0 else 'FAFBFC')
 text(s,41,y+6,330,20,risk,11.5);text(s,391,y+6,358,20,response,11.5)
text(s,32,450,730,22,'PROPOSED PILOT — NOT ACHIEVED: 1 society · 100 workers · 300 requests · target ≥90% completion.',10.7,BLUE,True)
text(s,32,473,728,22,'Production gate: governance, worker terms, pricing, legal obligations, data consent and provider access.',9.6,GRAY)
s.notes_slide.notes_text_frame.text='The pilot is proposed, not achieved. No speculative cost estimate or legal-clearance claim is made. Confirm the society actually supplies the selected household/community skills. Cash disputes require reconciliation and cannot trigger duplicate digital charging. No-show allegations must remain unverified until review.'

# 5 — audience, benefit categories, outcomes to measure.
s=base(5)
header(s,31,211,730,'Potential impact on the target audience',14)
audiences=[('Workers','Clearer terms, visible earnings, skill records and welfare guidance.'),('Customers','Discoverable cooperative workers, price visibility and traceable complaints.'),('Societies','Organized rosters, allocation tools and settlement visibility.'),('Federations','Authorized capacity and skill-gap summaries—not control over every job.')]
for i,(aud,desc) in enumerate(audiences):
 col=i%2;row=i//2;x=31+col*379;y=239+row*57
 text(s,x,y,347,18,aud,13,BLUE,True);text(s,x,y+19,345,33,desc,11.5)
header(s,31,359,730,'Benefits of the solution (social, economic, environmental, etc.)',13)
text(s,31,382,350,21,'Social: accessible, accountable participation.',11.8)
text(s,31,405,350,23,'Economic: aim for less idle time and clearer payments.',11.8)
text(s,410,382,350,23,'Environmental: test travel savings from nearby matching.',11.8)
text(s,410,405,350,23,'Governance: documented, approved cooperative decisions.',11.8)
shape(s,30,435,731,31,'EFF4F7')
text(s,40,439,711,27,'MEASURE: completion · utilization · assignment time · settlement delays\ncomplaint resolution · worker net earnings after reported costs',10.8,BLUE,True)
text(s,31,475,730,20,'Digitize the workforce. Keep the cooperative in control.',13.5,GREEN,True)
s.notes_slide.notes_text_frame.text='Benefits are hypotheses to evaluate in a pilot, not achieved impact. Workload-aware ranking alone does not prove fair allocations. Evaluate new-worker inclusion and avoid automatically ranking unrated workers last. Nearby matching may reduce travel, but the environmental effect must be measured.'

# 6 — concise official sources; do not invent PS/deployment URLs.
s=base(6)
header(s,31,216,730,'Details / Links of the reference and research work',15)
refs=[('01','SIH PS 26089 — organizer-provided statement','Confirmed portal URL required; title and theme supplied by the team.',None),('02','Ministry of Cooperation — National Cooperative Database','Society information and ecosystem planning; not a worker roster API.','https://www.cooperation.gov.in/en/national-cooperative-database'),('03','Haryana Labourfed — About Us','Documented society/district/state example—not a universal hierarchy.','https://labourfed.haryana.gov.in/about-us'),('04','International Cooperative Alliance — Cooperative Identity','Member ownership, democratic control and autonomy.','https://ica.coop/en/cooperatives/cooperative-identity'),('05','e-Shram — Official FAQs','Registration and welfare facilitation; scheme eligibility remains separate.','https://eshram.gov.in/faqs')]
for i,(n,title,desc,url) in enumerate(refs):
 y=252+i*37;text(s,32,y,25,22,n,13,BLUE,True);text(s,67,y,687,19,title,12,BLUE,True,url=url);text(s,67,y+18,687,18,desc,10,GRAY)
text(s,31,447,730,19,'Prototype links — URLs pending; add verified links before submission',11.5,BLUE,True)
for i,name in enumerate(['Household','Institution','Worker','Society','District','National']):
 x=31+i*124;text(s,x,473,118,18,name,11.5,INK,True)
s.notes_slide.notes_text_frame.text='Do not invent live prototype links. The supplied https://labourlink-customer.netlify.app was previously observed only as a landing page, not an end-to-end verified booking flow. This generated preview is not a claim of six separately deployed apps. NLCF About (https://labcofed.in/about-nlcf/) supports training, research and advocacy but not a mandatory national payment chain. ILO freedom of association: https://www.ilo.org/topics-and-sectors/freedom-association. Approved 2026 branding, registered team details, confirmed PS portal URL and final app URLs require team confirmation. The source template requests portal submission as PDF; the user requested this working PPTX only. The seventh instructions slide has intentionally been omitted.'

assert len(prs.slides)==6
path=OUT/'LABOURLINK-six-slide.pptx';prs.save(path)
(ROOT/'content-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
print(f'Created {path} · {len(prs.slides)} slides · {W} × {H} pt · {path.stat().st_size:,} bytes')